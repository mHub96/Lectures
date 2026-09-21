import os, sys
sys.stdout.reconfigure(encoding='utf-8')

from pptx import Presentation
import win32com.client

os.makedirs('verification_exports', exist_ok=True)

files = [
    'Ventricular Tumors_(Part 1).pptx',
    'Ventricular Tumors_(Part 2).pptx'
]

print("=== VERIFYING PRESENTATION STRUCTURE & XML ===")
for fname in files:
    prs = Presentation(fname)
    slides = list(prs.slides)
    print(f"\n{fname}: {len(slides)} slides")
    transitions_count = sum(1 for s in slides if '<p:transition' in s._element.xml)
    timings_count = sum(1 for s in slides if '<p:timing' in s._element.xml)
    images_count = sum(sum(1 for sh in s.shapes if sh.shape_type == 13) for s in slides)
    print(f"  Transitions: {transitions_count} / {len(slides)} slides")
    print(f"  Timings/Animations: {timings_count} / {len(slides)} slides")
    print(f"  Total Embedded Images: {images_count}")

print("\n=== EXPORTING REPRESENTATIVE SLIDES TO PNG ===")
ppt_app = win32com.client.Dispatch("PowerPoint.Application")

test_exports = [
    ('Ventricular Tumors_(Part 1).pptx', [1, 2, 10, 14, 15, 26, 32], 'part1'),
    ('Ventricular Tumors_(Part 2).pptx', [1, 3, 16, 18, 19, 23, 26], 'part2')
]

for ppt_file, slide_indices, prefix in test_exports:
    pres = ppt_app.Presentations.Open(os.path.abspath(ppt_file), WithWindow=False)
    for idx in slide_indices:
        out_png = os.path.abspath(f"verification_exports/{prefix}_slide{idx:02d}.png")
        pres.Slides(idx).Export(out_png, "PNG", 1920, 1080)
        print(f"Exported {out_png} ({os.path.getsize(out_png)//1024} KB)")
    pres.Close()

ppt_app.Quit()
print("\nAll verification exports completed successfully.")

