import sys
sys.stdout.reconfigure(encoding='utf-8')
from pptx import Presentation

prs = Presentation('Ventricular Tumors_(Part 1) (1).pptx')
for idx in [1, 2, 3, 4]:
    s = prs.slides[idx]
    print(f'=== Slide {idx+1} ===')
    for sh in s.shapes:
        txt = repr(sh.text[:50]) if sh.has_text_frame else ''
        print(f'  {sh.name} (type {sh.shape_type}): pos=({sh.left/914400:.2f}in, {sh.top/914400:.2f}in) size=({sh.width/914400:.2f}in x {sh.height/914400:.2f}in) text={txt}')

